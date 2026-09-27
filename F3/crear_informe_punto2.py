"""Actualiza la copia institucional con el informe completo de Sumativa 2."""
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Pt, RGBColor

RAIZ = Path(__file__).resolve().parents[1]
SALIDA = Path(r"C:\Magister\sumativo 2\informe_f3_grupo_5.docx")


def reemplazar(p, texto):
    if p.runs:
        p.runs[0].text = texto
        for r in p.runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        p.add_run(texto)


def limpiar_cuerpo(doc):
    for p in list(doc.paragraphs[17:]):
        p._element.getparent().remove(p._element)
    for t in list(doc.tables):
        t._element.getparent().remove(t._element)


def agregar_indice(doc):
    p = doc.add_paragraph(style="toc 1")
    r = p.add_run()
    ini = OxmlElement("w:fldChar"); ini.set(qn("w:fldCharType"), "begin")
    ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = ' TOC \\o "1-2" \\h \\z \\u '
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    r._r.extend([ini, ins, sep])
    entradas = (
        "Introducción\t3",
        "II. Diseño de soluciones algorítmicas eficientes\t3",
        "    A. Codificación funcional y arquitectura básica del script\t3",
        "    B. Preprocesamiento y transformación del dataset\t4",
        "    C. Validación técnica y verificación del código\t5",
        "    D. Eficiencia y optimización\t6",
        "    E. Diseño estructurado del código\t7",
        "III. Implementación de código modular y robusto\t8",
        "    A. Programación orientada a objetos\t8",
        "    B. Documentación de arquitectura y funcionalidad\t9",
        "IV. Repositorio GitHub F3\t10",
        "V. Notebooks ejecutables F3\t11",
        "VI. Bibliografía\t12",
        "Anexo A. Matriz de trazabilidad\t13",
    )
    for i, entrada in enumerate(entradas):
        if i:
            r._r.append(OxmlElement("w:br"))
        txt = OxmlElement("w:t"); txt.set(qn("xml:space"), "preserve"); txt.text = entrada; r._r.append(txt)
    fin = OxmlElement("w:fldChar"); fin.set(qn("w:fldCharType"), "end"); r._r.append(fin)


def titulo(doc, texto, nivel):
    doc.add_paragraph(texto, style="Título 11" if nivel == 1 else "Título 21")


def parrafo(doc, texto):
    p = doc.add_paragraph(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(6)


def vineta(doc, texto):
    p = doc.add_paragraph(style="Normal")
    p.add_run("• ").bold = True
    p.add_run(texto)
    p.paragraph_format.left_indent = Pt(18)
    p.paragraph_format.first_line_indent = Pt(-10)
    p.paragraph_format.space_after = Pt(3)


def codigo(doc, texto):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(24)
    p.paragraph_format.right_indent = Pt(12)
    p.paragraph_format.space_after = Pt(8)
    shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F2F5F8")
    p._p.get_or_add_pPr().append(shd)
    r = p.add_run(texto); r.font.name = "Courier New"; r.font.size = Pt(9)


def referencia(doc, texto, url=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Pt(25)
    p.paragraph_format.first_line_indent = Pt(-25)
    p.paragraph_format.space_after = Pt(6)
    p.add_run(texto)
    if url:
        p.add_run(" ")
        rid = doc.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
        enlace = OxmlElement("w:hyperlink"); enlace.set(qn("r:id"), rid)
        run = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1")
        sub = OxmlElement("w:u"); sub.set(qn("w:val"), "single")
        rpr.extend([color, sub]); run.append(rpr)
        nodo = OxmlElement("w:t"); nodo.text = url; run.append(nodo)
        enlace.append(run); p._p.append(enlace)


def tabla(doc, leyenda, encabezados, filas):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4)
    rr = p.add_run(leyenda); rr.bold = True; rr.italic = True
    t = doc.add_table(rows=1, cols=len(encabezados)); t.style = "Grid Table 1 Light Accent 1"
    for i, valor in enumerate(encabezados):
        c = t.rows[0].cells[i]; c.text = valor; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "24364B"); c._tc.get_or_add_tcPr().append(shd)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in c.paragraphs[0].runs:
            run.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(9)
    for n, fila in enumerate(filas, 1):
        cells = t.add_row().cells
        for i, valor in enumerate(fila):
            cells[i].text = str(valor); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
            if n % 2 == 0:
                shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), "F2F5F8"); cells[i]._tc.get_or_add_tcPr().append(shd)
            for run in cells[i].paragraphs[0].runs:
                run.font.size = Pt(9)
    doc.add_paragraph()


def construir():
    doc = Document(SALIDA); limpiar_cuerpo(doc); p = doc.paragraphs
    portada = {1:"MCDI500 • PROGRAMACIÓN PARA LA CIENCIA DE DATOS",2:"MAGÍSTER EN CIENCIA DE DATOS E INTELIGENCIA ARTIFICIAL",3:"UNIVERSIDAD ANDRÉS BELLO",5:"Análisis de ofertas en licitaciones públicas del sector Salud",6:"Avance Fase 3 – Semana 2",7:"Diseño algorítmico eficiente y programación orientada a objetos",9:"Grupo 5",10:"Víctor Bravo Barrera",11:"Nayadeth Garrido Ibáñez",12:"Mauricio Cid",14:"Docente: Omar Salinas Silva",15:"Fecha: 27 de septiembre de 2026"}
    for i, texto in portada.items():
        reemplazar(p[i], texto)

    idx = doc.add_paragraph("Índice de contenidos", style="Title"); idx.alignment = WD_ALIGN_PARAGRAPH.LEFT
    agregar_indice(doc); doc.add_page_break()

    titulo(doc,"Introducción",1)
    parrafo(doc,"Este informe presenta el avance de la Fase 3 del proyecto de análisis de ofertas en licitaciones públicas del sector Salud. El trabajo conserva la preparación reproducible construida en F1 y F2 y la amplía con implementaciones algorítmicas comparables, un núcleo orientado a objetos para lectura, limpieza y validación, pruebas automatizadas y mediciones de tiempo y memoria. La organización responde a la Evaluación Formativa 3 y a la pauta de la Evaluación Sumativa 2 (Universidad Andrés Bello, 2026a, 2026b).")
    parrafo(doc,"El resultado principal es la selección de la implementación agrupada con pandas para el conjunto completo de 44.226 registros. Esta alternativa obtuvo la menor mediana en las dos variables de agrupación estudiadas, sin cambiar la unidad de análisis ni los resultados de F2. La decisión se limita a este dataset y se apoya en equivalencia funcional, pruebas de error y mediciones reproducibles.")

    titulo(doc,"II. Diseño de soluciones algorítmicas eficientes",1)
    parrafo(doc,"Esta etapa integra el procesamiento desarrollado en F2 con tres estrategias para calcular la proporción de ofertas ganadoras por tamaño de proveedor y tipo de licitación. El flujo se evaluó sobre las 44.226 filas del reporte del sector Salud de marzo de 2026. Todas las variantes entregan los mismos resultados y se comparan mediante pruebas, mediciones de tiempo y estimaciones de memoria reproducibles.")

    titulo(doc,"A. Codificación funcional y arquitectura básica del script",2)
    parrafo(doc,"El código separa lectura, limpieza, validación y análisis. Las interfaces de F1 y F2 se conservan para no romper los notebooks anteriores. src/analisis.py mantiene tabla_proporciones como referencia y añade las variantes iterativa, recursiva y agrupada. Las tres reciben un DataFrame, una variable de agrupación y el estado administrativo, y devuelven una tabla con conteos, denominadores y porcentajes.")
    parrafo(doc,"La versión iterativa mantiene contadores por grupo. La recursiva divide intervalos de índices y combina diccionarios parciales, con un bloque base configurable de 256 registros. La versión agrupada utiliza operaciones groupby de pandas, diseñadas para separar, aplicar y combinar datos por categorías (The pandas development team, 2026). Las funciones comparten la preparación y la construcción de salida, lo que evita repetir validaciones y reglas de negocio.")

    titulo(doc,"B. Preprocesamiento y transformación del dataset",2)
    parrafo(doc,"El pipeline reutiliza el dataset procesado en F2 y conserva las 44.226 observaciones. Comprueba las columnas obligatorias, elimina tres variables completamente vacías, convierte diez variables de fecha, neutraliza 7.613 fechas anómalas del año 1900 y elimina espacios sobrantes en las categorías. También genera oferta_ganadora, licitacion_adjudicada y plazo_cierre_dias. Al eliminar y crear tres variables, el resultado conserva 74 columnas.")
    tabla(doc,"Tabla 1  Transformaciones reutilizadas en F3",("Etapa","Tratamiento","Justificación"),(
        ("Esquema","Cinco columnas obligatorias.","Evita procesar un archivo incompleto."),("Columnas vacías","Exclusión de tres variables 100 % vacías.","No contienen información utilizable."),("Fechas","Conversión a datetime y NaT para 1900.","Calcula plazos sin fechas ficticias."),("Texto","Eliminación de espacios en los extremos.","Unifica etiquetas sin cambiar su significado."),("Variables derivadas","Dos indicadores y un plazo en días.","Facilita las validaciones y el análisis.")))
    parrafo(doc,"No se aplicó escalamiento porque esta etapa compara frecuencias y proporciones, no entrena un modelo. TamanoProveedor conserva la categoría NoClasificado para no asignar artificialmente esas ofertas a otro tamaño empresarial.")

    titulo(doc,"C. Validación técnica y verificación del código",2)
    parrafo(doc,"La verificación combina doce pruebas del núcleo POO con ocho pruebas algorítmicas. Estas incluyen un resultado manual conocido, equivalencia con F2, grupos sin resultados válidos, valores ausentes, categorías reservadas, entradas vacías, columnas ausentes, índices repetidos, bloques recursivos inválidos, particiones aleatorias y el dataset real en sus dos agrupaciones. También se comprueba que una lista vacía de reglas produzca un error y que modificar el resumen del limpiador no cambie su estado interno. Las 20 pruebas finalizaron correctamente.")
    tabla(doc,"Tabla 2  Cobertura de verificación",("Caso","Comprobación","Resultado"),(
        ("Normal","Lectura, limpieza, validación y conteo.","Correcto"),("Equivalencia","Referencia, iterativa, recursiva y agrupada.","Correcto"),("Límite","Vacíos, NA, otros estados e índices repetidos.","Correcto"),("Excepción","Columnas, fechas, configuración y reglas inválidas.","Controlada"),("Encapsulamiento","Cambio del resumen externo del limpiador.","Estado intacto"),("Total","Doce pruebas POO y ocho algorítmicas.","20 de 20")))
    parrafo(doc,"La trazabilidad queda en F3/test_nucleo_poo.py, F3/test_algoritmos.py, F3/F3_Algoritmos.ipynb y evidencias/F3_algoritmos/. El arnés valida la equivalencia antes de registrar cada medición.")

    titulo(doc,"D. Eficiencia y optimización",2)
    parrafo(doc,"La medición utilizó timeit con siete rondas de tres ejecuciones, orden alternado, calentamiento previo y semilla 2026. Se informa la mediana por llamada. Las muestras anidadas contienen 100, 1.000, 10.000 y 44.226 filas. El tiempo incluye filtrado, preparación, conteo y formato de salida, pero excluye la lectura y limpieza del CSV. La repetición controlada reduce errores comunes de temporización (Python Software Foundation, 2026a).")
    tabla(doc,"Tabla 3  Mediana en el conjunto completo en milisegundos",("Agrupación","Referencia F2","Iterativa","Recursiva","Agrupada"),(("TamanoProveedor","66,922","41,857","41,691","38,554"),("TipoLicitacion","68,251","40,716","40,528","39,760")))
    tabla(doc,"Tabla 4  Complejidad y memoria observada",("Implementación","Complejidad temporal","Pico trazado"),(("Referencia F2","O(n)","40,81 MiB"),("Iterativa","O(n)","26,07 MiB"),("Recursiva divide y vencerás","O(n + S), pila O(log L)","26,07 MiB"),("Agrupada con pandas","O(n)","26,07 MiB")))
    parrafo(doc,"Se selecciona tabla_proporciones_agrupada para F3 porque obtuvo la menor mediana en ambas variables y un pico trazado similar al de las otras alternativas nuevas. Frente a la referencia F2, la mejora observada fue de aproximadamente 1,74 veces por tamaño de proveedor y 1,72 veces por tipo de licitación. La ventaja sobre las versiones iterativa y recursiva es menor, por lo que la elección se limita al volumen y las categorías de este estudio.")
    parrafo(doc,"Los picos se midieron con tracemalloc y representan asignaciones rastreadas por Python, no toda la memoria nativa de pandas o NumPy (Python Software Foundation, 2026b). Por ello se interpretan como una comparación controlada dentro del mismo entorno y no como el consumo total del proceso.")

    titulo(doc,"E. Diseño estructurado del código",2)
    parrafo(doc,"El flujo se organiza en cinco pasos: leer, limpiar, validar, calcular y medir. Cada paso recibe una entrada definida y devuelve un resultado comprobable. src/datos.py y src/pipeline.py concentran la preparación; src/validacion.py aplica las reglas; src/analisis.py contiene las variantes comparables; y F3 separa pruebas, notebook y mediciones. Esta distribución evita mezclar la presentación con la lógica de procesamiento.")
    tabla(doc,"Tabla 5  Flujo y responsabilidades",("Paso","Módulo","Comprobación antes de avanzar"),(("1  Leer","src/datos.py","Archivo, registros y esquema mínimo."),("2  Limpiar","src/pipeline.py","Fechas, texto y variables derivadas."),("3  Validar","src/validacion.py","Seis reglas de calidad superadas."),("4  Calcular","src/analisis.py","Equivalencia de las implementaciones."),("5  Medir","F3/medir_algoritmos.py","Tiempos, memoria, entorno y hashes.")))
    parrafo(doc,"La recursividad divide el intervalo hasta alcanzar bloques pequeños y luego combina resultados parciales. La profundidad es logarítmica respecto del número de bloques. Aunque no entrega una ventaja sostenida frente al recorrido iterativo, demuestra descomposición, terminación y combinación correcta sobre un problema real.")

    doc.add_page_break()
    titulo(doc,"III. Implementación de código modular y robusto",1)
    titulo(doc,"A. Programación orientada a objetos",2)
    parrafo(doc,"La solución orientada a objetos distribuye responsabilidades entre contratos, lectores, limpiadores y reglas de validación. Esta separación evita una clase única que concentre todo el proceso y permite ampliar el sistema sin modificar componentes que ya funcionan, lo que favorece un diseño sostenible y extensible (Phillips, 2021).")
    parrafo(doc,"Herencia. ReglaValidacion define el contrato común de las comprobaciones. ReglaFechas, ReglaColumnasObligatorias y ReglaVariablesDerivadas heredan de ella e implementan su propio nombre y método validar. LectorCSV aplica de forma equivalente el contrato abstracto definido por LectorDatos.")
    parrafo(doc,"Polimorfismo. ValidadorDatasetProcesado recorre una colección de reglas y llama al mismo método validar, sin conocer el detalle de cada comprobación. Una regla nueva puede integrarse si hereda de ReglaValidacion; las pruebas lo demuestran con ReglaSiempreValida, sin cambiar el código del validador.")
    parrafo(doc,"Encapsulamiento. LimpiadorLicitaciones guarda la trazabilidad en _ultima_ejecucion y la expone mediante una propiedad que devuelve una copia. Si el código externo modifica el resumen recibido, el estado interno permanece intacto. La prueba cambia el número de filas de la copia a 999 y confirma que el objeto conserva el valor real de 2 en el caso controlado.")
    tabla(doc,"Tabla 6  Clases principales del núcleo F3",("Clase","Responsabilidad","Relación"),(
        ("ContratoEsquema","Define las columnas mínimas y valida registros.","Usado por los lectores."),
        ("LectorDatos","Declara la interfaz abstracta de lectura.","Base de LectorCSV."),
        ("LectorCSV","Lee el CSV y aplica el contrato.","Recibe un ContratoEsquema."),
        ("LimpiadorLicitaciones","Coordina las transformaciones y registra la ejecución.","Usa funciones de preprocesamiento."),
        ("ReglaValidacion","Declara la interfaz de cada regla.","Base de seis reglas concretas."),
        ("ValidadorDatasetProcesado","Ejecuta reglas intercambiables.","Recibe una colección de reglas.")))
    codigo(doc,"# Copia defensiva del estado interno\n@property\ndef ultima_ejecucion(self):\n    if self._ultima_ejecucion is None:\n        return None\n    return self._ultima_ejecucion.copy()")

    titulo(doc,"B. Documentación de arquitectura y funcionalidad del código",2)
    parrafo(doc,"Los métodos públicos incluyen docstrings breves y mensajes de error que indican qué dato debe corregirse. Las decisiones se registran en docs/DECISIONES_TECNICAS.md y docs/F3_APORTE_ALGORITMOS.md. El primer documento explica las alternativas evaluadas y los motivos de descarte; el segundo registra la unidad de análisis, la complejidad, el método de medición, los límites y la elección final.")
    parrafo(doc,"La compatibilidad se mantiene mediante src/proyecto.py, que reexporta funciones y clases desde módulos especializados. Así, los notebooks anteriores conservan sus importaciones mientras cada definición mantiene una sola implementación. El historial documenta la evolución desde funciones concentradas en F2 hacia componentes orientados a objetos y separados por responsabilidad en F3.")
    codigo(doc,"# Diferencia explícita entre ausencia y colección vacía\nself._reglas = tuple(\n    reglas_predeterminadas if reglas is None else reglas\n)\nif not self._reglas:\n    raise ValueError('El validador requiere al menos una regla.')")

    doc.add_page_break()
    titulo(doc,"IV. Repositorio GitHub F3",1)
    parrafo(doc,"El repositorio conserva las carpetas F1, F2 y F3, junto con src, data, docs y evidencias. F3 contiene el notebook, las pruebas, el script de medición, el verificador y los archivos de ejemplo. README.md y F3/README.md describen dependencias, estructura y comandos de ejecución. Los datos procesados no se duplican en Git porque se regeneran desde el archivo original; los hashes SHA-256 permiten comprobar la versión utilizada.")
    tabla(doc,"Tabla 7  Evolución registrada en Git",("Fecha","Cambio","Evidencia"),(
        ("23-09-2026","Incorporación inicial del núcleo POO.","Commit 8087c45"),
        ("24-09-2026","Separación por responsabilidad con compatibilidad F1/F2.","Commit 189f251"),
        ("26-09-2026","Recursividad, mediciones y equivalencia algorítmica.","Commit 48e62d3"),
        ("27-09-2026","Cierre de limpieza, validación y pruebas de error.","Commit d90111e"),
        ("27-09-2026","Corrección del código de curso a MCDI.","Commit ab661e4")))
    parrafo(doc,"Las contribuciones individuales visibles en el historial se distribuyen de la siguiente manera:")
    vineta(doc,"Víctor Bravo Barrera: separación modular, algoritmos, recursividad, mediciones y notebook F3.")
    vineta(doc,"Nayadeth Garrido Ibáñez: núcleo POO, limpieza, validación, casos de error y correcciones documentales.")
    vineta(doc,"Mauricio Cid: decisiones técnicas, documentación del flujo reproducible y antecedentes visuales del proyecto.")
    parrafo(doc,"El repositorio está disponible en https://github.com/vbravo29/sumativo-1. Su estructura coincide con los módulos citados en este informe y permite rastrear cada resultado hacia código, prueba o evidencia.")

    titulo(doc,"V. Notebooks ejecutables F3",1)
    parrafo(doc,"F3/F3_Algoritmos.ipynb integra el dataset procesado de F2, el núcleo POO, las implementaciones de referencia, iterativa, recursiva y agrupada, los casos de prueba y las mediciones. Su organización separa entorno, lectura, limpieza, calidad, ejemplo manual, recursividad, pruebas, resultados, tiempos, memoria y decisión final. Cada celda de salida presenta como máximo una tabla para facilitar la revisión.")
    tabla(doc,"Tabla 8  Reproducibilidad del notebook",("Elemento","Implementación","Evidencia"),(
        ("Entorno","requirements.txt y kernel del proyecto.","Versiones registradas."),
        ("Datos","CSV original y pipeline F2/F3.","Dimensiones y SHA-256."),
        ("Pruebas","unittest para algoritmos y clases.","20 pruebas aprobadas."),
        ("Medición","timeit y tracemalloc por separado.","CSV y JSON en evidencias."),
        ("Ejecución","F3/verificar_algoritmos.py en kernel nuevo.","24 celdas y cero errores en el registro base.")))
    parrafo(doc,"Para reproducir el avance se crea el entorno, se instalan las dependencias, se ubica el CSV en data/raw, se ejecutan las pruebas, se regeneran las mediciones cuando cambia el código y finalmente se ejecuta el verificador del notebook. Los hashes evitan presentar resultados antiguos después de modificar src/analisis.py o el script de medición.")
    codigo(doc,'.\\.venv\\Scripts\\python.exe -m unittest discover -s F3 -p "test_*.py" -v\n.\\.venv\\Scripts\\python.exe F3\\medir_algoritmos.py\n.\\.venv\\Scripts\\python.exe F3\\verificar_algoritmos.py')

    doc.add_page_break()
    titulo(doc,"VI. Bibliografía",1)
    referencia(doc,"McKinney, W. (2022). Python for data analysis: Data wrangling with pandas, NumPy, and Jupyter (3.ª ed.). O’Reilly Media.","https://wesmckinney.com/book/")
    referencia(doc,"Phillips, D. (2021). Python 3 object-oriented programming (4.ª ed.). Packt Publishing.","https://www.packtpub.com/en-us/product/python-3-object-oriented-programming-9781801077262")
    referencia(doc,"Python Software Foundation. (2026a). timeit — Measure execution time of small code snippets. Python 3 documentation.","https://docs.python.org/3/library/timeit.html")
    referencia(doc,"Python Software Foundation. (2026b). tracemalloc — Trace memory allocations. Python 3 documentation.","https://docs.python.org/3/library/tracemalloc.html")
    referencia(doc,"The pandas development team. (2026). pandas.DataFrame.groupby. pandas documentation.","https://pandas.pydata.org/docs/reference/api/pandas.DataFrame.groupby.html")
    referencia(doc,"Universidad Andrés Bello. (2026a). Evaluación Formativa 3: núcleo algorítmico, eficiencia y programación orientada a objetos [Material docente del curso MCDI500].")
    referencia(doc,"Universidad Andrés Bello. (2026b). Evaluación Sumativa 2: avance del proyecto Fase 3, Semana 2 [Pauta y rúbrica del curso MCDI500].")

    titulo(doc,"Anexo A. Matriz de trazabilidad",1)
    tabla(doc,"Tabla A1  Relación entre requisitos y evidencia",("Requisito","Archivo principal","Evidencia"),(
        ("Pipeline de limpieza","src/pipeline.py y src/preprocesamiento.py","44.226 filas y 74 columnas."),
        ("Validación","src/validacion.py","Seis reglas y errores explícitos."),
        ("POO","src/datos.py, pipeline.py y validacion.py","Herencia, polimorfismo y encapsulamiento."),
        ("Algoritmos","src/analisis.py","Cuatro implementaciones equivalentes."),
        ("Pruebas","F3/test_nucleo_poo.py y test_algoritmos.py","20 pruebas aprobadas."),
        ("Rendimiento","F3/medir_algoritmos.py","Tiempos, memoria, versiones y hashes."),
        ("Notebook","F3/F3_Algoritmos.ipynb","Ejecución estructurada y trazable.")))

    for nombre in ("Title","Título 11","Título 21","toc 1","toc 2"):
        if nombre in doc.styles:
            doc.styles[nombre].font.color.rgb = RGBColor(0,0,0)
    upd = OxmlElement("w:updateFields"); upd.set(qn("w:val"), "true"); doc.settings._element.append(upd)
    doc.core_properties.title = "Análisis de ofertas en licitaciones públicas del sector Salud Fase 3"
    doc.save(SALIDA)
    print(SALIDA)


if __name__ == "__main__":
    construir()
